"""Lossless source adapters. Dictionary matches are evidence, never sense assignments."""

import io
import re
import urllib.parse
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from pathlib import Path

from .common import NCAT_NOTICE, aliases, digest, record
from .network import CONCISED_VERSION, CONCISED_URL, MS_URL


def linguipedia_records(snapshot):
    if not snapshot.get("complete"):
        raise ValueError("Cannot import an incomplete Linguipedia snapshot")
    items = snapshot["items"]
    if len(items) != snapshot["expected_total"] or len({r["id"] for r in items}) != len(items):
        raise ValueError("Linguipedia snapshot count/ID mismatch")
    if digest(items) != snapshot["sha256"]:
        raise ValueError("Linguipedia snapshot checksum mismatch")
    result = [record("linguipedia", r["id"], tw=aliases(r.get("tw_word") or ""),
        cn=aliases(r.get("cn_word") or ""), domain=r["category"], relation=r["type"],
        source_forms={"tw": r.get("tw_word") or "", "cn": r.get("cn_word") or ""},
        original=r, version=snapshot["sha256"],
        url="https://www.chinese-linguipedia.org/search_difference.html?word=" + urllib.parse.quote(r.get("tw_word") or r.get("cn_word") or "")) for r in items]
    return result, {"complete": True, "expected": snapshot["expected_total"], "imported": len(result),
                    "pages": snapshot["completed_pages"], "sha256": snapshot["sha256"],
                    "retrieved_at": snapshot["fetched_at"], "relations": dict(Counter(r["relation"] for r in result)),
                    "notice": NCAT_NOTICE}


def concised_records(path, version=CONCISED_VERSION):
    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError("Install import dependencies: python3 -m pip install -r requirements-import.txt") from e
    path = Path(path)
    with zipfile.ZipFile(path) as archive:
        files = [n for n in archive.namelist() if n.endswith(".xlsx") and not n.startswith("__MACOSX/")]
        if len(files) != 1:
            raise ValueError("Expected exactly one Concised XLSX workbook")
        workbook = openpyxl.load_workbook(io.BytesIO(archive.read(files[0])), read_only=True, data_only=True)
    records, stats, seen = [], [], set()
    for sheet in workbook:
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        if not {"字詞名", "字詞號", "注音一式", "釋義", "多音排序"}.issubset(headers):
            raise ValueError(f"Unrecognized Concised worksheet: {sheet.title}")
        processed, empty = 0, 0
        for row in rows:
            if not any(v is not None and str(v).strip() for v in row):
                empty += 1
                continue
            original = dict(zip(headers, row))
            title, identity = original["字詞名"], original["字詞號"]
            if not isinstance(title, str) or not isinstance(identity, str) or not identity:
                raise ValueError(f"Invalid Concised entry in {sheet.title}: {identity!r}")
            key = identity + ":" + str(original["多音排序"] or 0)
            if key in seen:
                raise ValueError(f"Duplicate Concised entry identity: {key}")
            seen.add(key)
            meaning = original["釋義"] or ""
            # Keep the original Excel cell verbatim, including any OOXML escape.
            display = meaning.replace("_x000D_", "\r")
            pronunciation = [{"bopomofo": original["注音一式"] or "", "pinyin": original.get("漢語拼音") or "",
                              "order": original["多音排序"], "variant_bopomofo": original.get("變體注音") or "",
                              "variant_pinyin": original.get("變體漢語拼音") or ""}]
            records.append(record("moe-concised", key, tw=[title.strip()], kind="dictionary",
                definitions=[{"text": meaning, "display": display}] if meaning else [],
                pronunciations=pronunciation, original=original, domain="一般用語", version=version,
                url="https://dict.concised.moe.edu.tw/search.jsp?word=" + urllib.parse.quote(title.strip())))
            processed += 1
        if sheet.max_row is not None and processed + empty + 1 != sheet.max_row:
            raise ValueError(f"Concised worksheet row count mismatch: {sheet.title}")
        stats.append({"sheet": sheet.title, "total_rows": processed + empty + 1, "header_rows": 1,
                      "imported": processed, "blank_rows": empty})
    workbook.close()
    return records, {"complete": True, "version": version, "url": CONCISED_URL,
        "sha256": digest(path.read_bytes()), "worksheets": stats, "imported": len(records),
        "unique_headwords": len({r["tw"][0] for r in records}),
        "notice": "中華民國教育部。《國語辭典簡編本》（版本編號：" + version + "）。CC BY-ND 3.0 TW；完整使用說明見 raw/concised-usage.pdf。"}


def microsoft_records(path, manifest):
    path = Path(path)
    if digest(path.read_bytes()) != manifest["sha256"]:
        raise ValueError("Microsoft TBX checksum mismatch")
    root = ET.parse(path).getroot()
    records, excluded, unknown = [], [], []
    count = 0
    language = "{http://www.w3.org/XML/1998/namespace}lang"
    for entry in root.iter("termEntry"):
        count += 1
        identity = entry.get("id")
        if not identity:
            raise ValueError("Microsoft entry missing ID")
        en, tw, definitions, original_terms = [], [], [], []
        for lang in entry.findall("langSet"):
            locale = lang.get(language)
            if locale == "en-US":
                en.extend((t.text or "") for t in lang.iter("term"))
                definitions.extend({"text": "".join(d.itertext())} for d in lang.iter("descrip") if d.get("type") == "definition")
            if locale not in ("zh-Hant", "zh-TW"):
                continue
            for group in lang.findall("ntig") + lang.findall("tig"):
                term = group.find(".//term")
                if term is None:
                    raise ValueError(f"Missing term in Microsoft entry {identity}")
                notes = [{"type": n.get("type"), "text": "".join(n.itertext())} for n in group.iter("termNote")]
                geography = [n["text"] for n in notes if n["type"] == "geographicalUsage"]
                regions = set(re.split(r"[,;\s]+", " ".join(geography).upper())) - {""}
                item = {"term": term.text or "", "locale": locale, "notes": notes, "id": term.get("id")}
                original_terms.append(item)
                if "TWN" in regions or locale == "zh-TW":
                    tw.append(term.text or "")
                elif regions:
                    excluded.append({"entry_id": identity, **item})
                else:
                    unknown.append({"entry_id": identity, **item})
        if tw and en:
            records.append(record("microsoft", identity, en=list(dict.fromkeys(en)), tw=list(dict.fromkeys(tw)),
                definitions=definitions, kind="technical", domain="電腦資訊", version=manifest["sha256"],
                part_of_speech=" / ".join(sorted({n["text"] for t in original_terms for n in t["notes"] if n["type"] == "partOfSpeech"})),
                geographical_usage=sorted({n["text"] for t in original_terms for n in t["notes"] if n["type"] == "geographicalUsage"}),
                url="https://learn.microsoft.com/en-us/globalization/reference/microsoft-terminology",
                original={"entry_id": identity, "terms": original_terms,
                          "xml": ET.tostring(entry, encoding="unicode")}))
    if count != manifest["expected_entries"]:
        raise ValueError("Microsoft entry count mismatch")
    if len({r["id"] for r in records}) != len(records):
        raise ValueError("Duplicate Microsoft source IDs")
    return records, {"complete": True, "url": MS_URL, "sha256": manifest["sha256"],
        "expected_entries": count, "processed_entries": count, "imported": len(records),
        "other_region_terms": len(excluded), "unknown_region_terms": len(unknown),
        "region_policy": "TWN or explicit zh-TW; unknown regions are quarantined, not treated as Taiwan",
        "last_modified": manifest["last_modified"]}, {"other_regions": excluded, "unknown_regions": unknown}
